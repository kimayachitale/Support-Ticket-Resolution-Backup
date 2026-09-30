from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from database import get_history
from datetime import datetime
import io

def export_excel():

    tickets = get_history()

    wb = Workbook()
    ws = wb.active
    ws.title = "Railway Tickets"

    # ===== Title =====
    ws.merge_cells("A1:E1")
    ws["A1"] = "Railway Support Ticket Resolution Platform"
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    ws["A1"].alignment = Alignment(horizontal="center")

    # ===== Date =====
    ws.merge_cells("A2:E2")
    ws["A2"] = "Export Date : " + datetime.now().strftime("%d-%m-%Y %H:%M")
    ws["A2"].font = Font(bold=True)
    ws["A2"].alignment = Alignment(horizontal="center")

    # ===== Header =====
    headers = ["ID", "Complaint", "Category", "PNR Number", "Decision"]

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="4F81BD")
        cell.alignment = Alignment(horizontal="center")

    thin = Side(style="thin")

    row = 5

    for ticket in tickets:

        ws.cell(row=row, column=1).value = ticket[0]
        ws.cell(row=row, column=2).value = ticket[1]
        ws.cell(row=row, column=3).value = ticket[2]
        ws.cell(row=row, column=4).value = ticket[3]
        ws.cell(row=row, column=5).value = ticket[4]

        for col in range(1, 6):
            cell = ws.cell(row=row, column=col)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

        row += 1

    # ===== Column Width =====
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 25
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 20

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="Railway_Tickets.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )